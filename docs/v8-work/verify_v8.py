from pathlib import Path
from zipfile import ZipFile
from collections import Counter
from lxml import etree as E
from docx import Document
from PIL import Image,ImageDraw
import pdfplumber, hashlib,json,re,importlib.util

WORK=Path(__file__).resolve().parent
ROOT=WORK.parents[1]
FINAL=WORK/'word-final.docx'
SOURCE=WORK/'source-v7.docx'
NS={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main','w14':'http://schemas.microsoft.com/office/word/2010/wordml'}
sha=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()
norm=lambda t:re.sub(r'\s+','',t)
native=json.loads((WORK/'word-native-check.json').read_text(encoding='utf-8-sig'))
changes=json.loads((WORK/'applied-changes.json').read_text(encoding='utf-8'))['changes']
final=Document(FINAL)
source=Document(SOURCE)

helper=Path('C:/Users/a3185/.codex/plugins/cache/openai-primary-runtime/documents/26.1020.11127/skills/documents/render_docx.py')
spec=importlib.util.spec_from_file_location('documents_renderer',helper)
renderer=importlib.util.module_from_spec(spec);spec.loader.exec_module(renderer)
# runtime.json confirms no bundled LibreOffice on Windows; use Word's native
# PDF as conversion result and keep the skill's Poppler rasterization unchanged.
renderer.convert_to_pdf=lambda *a,**kw:(str(WORK/'native-final.pdf'),'Native Word export; runtime LibreOffice version is null')
rendered=renderer.rasterize(str(FINAL),str(WORK/'render-final'),144,False,False)

with ZipFile(SOURCE) as a,ZipFile(FINAL) as b:
    assert b.testzip() is None
    ma={n:hashlib.sha256(a.read(n)).hexdigest() for n in a.namelist() if n.startswith('word/media/')}
    mb={n:hashlib.sha256(b.read(n)).hexdigest() for n in b.namelist() if n.startswith('word/media/')}
    assert Counter(ma.values())==Counter(mb.values())
    sx=E.fromstring(a.read('word/document.xml'));fx=E.fromstring(b.read('word/document.xml'))
    txt=lambda p:''.join(p.xpath('.//w:t/text()',namespaces=NS))
    new_by_id={p.get('{'+NS['w14']+'}paraId'):txt(p) for p in fx.xpath('.//w:p',namespaces=NS)}
    modified={c['paraId'] for c in changes}
    preserved=0
    for p in sx.xpath('.//w:p',namespaces=NS):
        pid=p.get('{'+NS['w14']+'}paraId')
        st=p.xpath('./w:pPr/w:pStyle/@w:val',namespaces=NS)
        if pid in modified or any(s.lower().startswith('toc') for s in st):continue
        assert pid in new_by_id,(pid,txt(p),'missing')
        assert new_by_id[pid]==txt(p),(pid,txt(p),new_by_id[pid])
        preserved+=1
    expected={c['paraId']:c['after'] for c in changes}
    assert all(new_by_id[pid]==val for pid,val in expected.items())
    footerparts=[n for n in b.namelist() if re.match(r'word/footer\d+\.xml$',n)]
    assert any(re.search(r'\bPAGE\b',b.read(n).decode('utf-8')) for n in footerparts)
    assert fx.xpath('.//w:instrText[contains(text(),"TOC")]',namespaces=NS)

code_pattern=r'^(?:FROM |WORKDIR |COPY |RUN |USER |ENV |EXPOSE |ENTRYPOINT |docker |curl\.exe |java |\.\\scripts\\|sha256:|Get-Item |  - |volumes:|name:|# )'
old_code=[p.text for p in source.paragraphs if re.match(code_pattern,p.text)]
new_code=[p.text for p in final.paragraphs if re.match(code_pattern,p.text)]
assert old_code==new_code
captions=[p.text for p in final.paragraphs if re.match(r'^图(?:\d+|A\d+)',p.text)]
assert len(captions)==len(final.inline_shapes)==31
assert len(final.tables)==7
all_text='\n'.join(p.text for p in final.paragraphs)+'\n'+'\n'.join(c.text for t in final.tables for r in t.rows for c in r.cells)
assert not any(s in all_text for s in ['待确认','TODO','占位符','构建缓存及其他基础镜像未作为','该次初始失败记录另行保存在','环境、资源和结果分别列示'])

with pdfplumber.open(WORK/'native-final.pdf') as pdf:
    assert len(pdf.pages)==native['pages']==len(rendered)
    texts=[p.extract_text() or '' for p in pdf.pages]
    pagechecks=[]
    for i,p in enumerate(pdf.pages):
        words=p.extract_words()
        clipped=[w['text'] for w in words if w['x0']< -1 or w['x1']>p.width+1 or w['top']< -1 or w['bottom']>p.height+1]
        assert not clipped,(i+1,clipped)
        assert len(texts[i])>25
        pagechecks.append({'page':i+1,'text_characters':len(texts[i]),'images':len(p.images),'out_of_page_words':clipped})
    for h in native['headings']:
        assert norm(h['heading']) in norm(texts[h['page']-1]),h
    toc=[]
    for block in native['tocText']:
        for line in block.split('\r'):
            bits=line.strip().split('\t')
            if len(bits)>1 and bits[-1].isdigit():toc.append({'heading':bits[0].strip(),'page':int(bits[-1])})
    assert len(toc)==len(native['headings'])==39
    for h in toc:
        assert any(norm(h['heading'])==norm(n['heading']) and h['page']==n['page'] for n in native['headings']),h
    assert len(native['figurePages'])==len(native['captionPages'])==31
    pairs=[]
    for page,caption in zip(native['figurePages'],native['captionPages']):
        pairs.append({'caption':caption['text'],'image_page':page,'caption_page':caption['page'],'same_page':page==caption['page']})
    assert all(p['same_page'] for p in pairs),[p for p in pairs if not p['same_page']]
    (WORK/'pdf-text.txt').write_text('\n\n'.join(f'PAGE {i+1}\n{t}' for i,t in enumerate(texts)),encoding='utf-8')

dest=WORK/'render-final'
for start in range(1,len(rendered)+1,9):
    end=min(start+8,len(rendered))
    sheet=Image.new('RGB',(1380,1935),'#d8d8d8');draw=ImageDraw.Draw(sheet)
    for offset,index in enumerate(range(start,end+1)):
        im=Image.open(dest/f'page-{index}.png');im.thumbnail((440,605))
        x=offset%3*460+10;y=offset//3*645+25
        sheet.paste(im,(x,y));draw.text((x,y-18),f'Page {index}',fill='black')
    sheet.save(dest/f'contact-{start}-{end}.jpg',quality=90)

baseline=json.loads((WORK/'baseline-sha256.json').read_text(encoding='utf-8'))
assert all(sha(ROOT/p)==h for p,h in baseline.items())
result={'sha256':sha(FINAL),'pages':native['pages'],'tables':7,'images':31,'toc_entries':39,'code_paragraphs_preserved':len(old_code),'unchanged_paragraphs_verified':preserved,'changed_paragraphs_verified':len(expected),'image_hashes_preserved':len(ma),'captions_same_page_as_images':pairs,'page_checks':pagechecks,'toc_entries_verified':toc,'baseline_files_unchanged':len(baseline),'renderer':'documents render_docx.py Poppler stage with native Word PDF adapter'}
(WORK/'final-checks.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
(WORK/'final-text.txt').write_text(all_text,encoding='utf-8')
print(json.dumps({k:result[k] for k in ['pages','tables','images','toc_entries','code_paragraphs_preserved','unchanged_paragraphs_verified','changed_paragraphs_verified','baseline_files_unchanged']},ensure_ascii=False))
