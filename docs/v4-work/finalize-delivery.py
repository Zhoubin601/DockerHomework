"""Validate and publish the verified v4 files without altering older releases."""

from collections import Counter
from pathlib import Path
from zipfile import ZipFile
import hashlib
import json
import re
import shutil

import fitz
from docx import Document
from lxml import etree

ROOT = Path(__file__).resolve().parents[2]
WORK = ROOT / 'docs/v4-work'
REPORT = ROOT / 'docs/v4-report-work'
DOCX = REPORT / '20261021-docker-lab-report-v4.docx'
PDF = REPORT / 'render-final/20261021-docker-lab-report-v4.pdf'
NS = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}


def sha(data):
    return hashlib.sha256(data).hexdigest()


def unchanged_baseline():
    baseline = json.loads((WORK / 'baseline-sha256.json').read_text())
    for name, record in baseline.items():
        data = (ROOT / name).read_bytes()
        assert len(data) == record['bytes'] and sha(data) == record['sha256'], name
    return len(baseline)


baseline_count = unchanged_baseline()
docx_hash = sha(DOCX.read_bytes())
for filename in ['visual-qa-pages1-15.md', 'visual-qa-pages16-29.md']:
    qa = (REPORT / filename).read_text()
    assert docx_hash in qa, f'Visual QA does not reference current DOCX: {filename}'

doc = Document(DOCX)
assert len(doc.inline_shapes) == 21 and len(doc.tables) == 8
paragraphs = '\n'.join(p.text for p in doc.paragraphs)
assert '图A1a 首次创建独立网络及数据卷' in paragraphs
assert '图A1b Linux引擎及网络和数据卷配置验证' in paragraphs
assert 'v4为报告与交付资料修订' in paragraphs
assert not re.search(r'TODO|TBD|PLACEHOLDER|待确认|未命名', paragraphs)

originals = {}
for directory in ['docs/docker-evidence', 'docs/v3-evidence']:
    for path in (ROOT / directory).glob('*.png'):
        originals.setdefault(sha(path.read_bytes()), []).append(str(path.relative_to(ROOT)))

with ZipFile(DOCX) as archive:
    assert archive.testzip() is None
    for name in archive.namelist():
        if name.endswith('.xml') or name.endswith('.rels'):
            etree.fromstring(archive.read(name))
    xml = etree.fromstring(archive.read('word/document.xml'))
    assert not xml.xpath('.//w:ins | .//w:del', namespaces=NS)
    bookmarks = xml.xpath('.//w:bookmarkStart/@w:name', namespaces=NS)
    assert len(bookmarks) == len(set(bookmarks))
    media = {
        name: originals.get(sha(archive.read(name)), [])
        for name in archive.namelist() if name.startswith('word/media/')
    }
    assert len(media) == 21 and all(media.values())
    assert any('docs/v3-evidence/01-resource-creation.png' in files for files in media.values())
    assert any('docs/v3-evidence/01-config-review.png' in files for files in media.values())

pdf = fitz.open(PDF)
assert len(pdf) == 29
normalize = lambda value: re.sub(r'\s+', '', value)
pages = [normalize(page.get_text()) for page in pdf]
toc = json.loads((REPORT / 'toc-headings.json').read_text())
assert len(toc) == 39
verified = []
for entry in toc:
    fields = xml.xpath('.//w:fldSimple[contains(@w:instr,$key)]', namespaces=NS,
                       key='PAGEREF ' + entry['bookmark'] + ' ')
    assert len(fields) == 1, entry
    cached = int(''.join(fields[0].xpath('.//w:t/text()', namespaces=NS)))
    actual = [index + 1 for index, text in enumerate(pages)
              if index >= 2 and normalize(entry['heading']) in text]
    assert actual == [cached], (entry, cached, actual)
    verified.append({**entry, 'page': cached})

source = WORK / '20261021-memory-game6085-v4.zip'
evidence = WORK / '20261021-docker-evidence-v4.zip'
with ZipFile(source) as archive:
    assert archive.testzip() is None
    for name in archive.namelist():
        assert archive.read(name) == (ROOT / 'Code' / name).read_bytes(), name
    jar_names = [name for name in archive.namelist() if name.endswith('.jar')]
    assert len(jar_names) == 1
    jar_hash = sha(archive.read(jar_names[0]))
with ZipFile(evidence) as archive:
    assert archive.testzip() is None
    manifest = json.loads(archive.read('manifest-sha256.json'))
    assert set(manifest) == set(archive.namelist()) - {'manifest-sha256.json'}
    for name, record in manifest.items():
        data = archive.read(name)
        assert sha(data) == record['sha256'] and len(data) == record['bytes'], name

result = json.loads((ROOT / 'docs/v3-evidence/results.json').read_text())
assert jar_hash == result['jarSha256'] == result['containerJarSha256']
validation = {
    'pages': len(pdf), 'tocEntries': len(toc), 'tocPagesMatchFinalRender': True,
    'bookmarksUnique': True, 'xmlAndZip': 'passed', 'figures': len(doc.inline_shapes),
    'tables': len(doc.tables), 'noDraftOrRevisionNodes': True,
    'all29PagesVisuallyInspected': True, 'mediaMatches': media,
    'rawAndOlderReleaseFilesUnchanged': baseline_count,
    'sha256': docx_hash,
}
(REPORT / 'structural-validation.json').write_text(json.dumps(validation, ensure_ascii=False, indent=2) + '\n')
(REPORT / 'toc-verified-pages.json').write_text(json.dumps(verified, ensure_ascii=False, indent=2) + '\n')

destination = ROOT / 'output/v4'
destination.mkdir(exist_ok=True)
pairs = [(DOCX, destination / DOCX.name), (source, destination / source.name),
         (evidence, destination / evidence.name)]
for kind in ['desktop', 'mobile']:
    pairs.append((ROOT / f'output/v3/20261021-memory-game6085-{kind}-v3.png',
                  destination / f'20261021-memory-game6085-{kind}-v4.png'))
for origin, target in pairs:
    if target.exists():
        assert target.read_bytes() == origin.read_bytes(), f'Refusing to replace existing file: {target}'
    else:
        shutil.copyfile(origin, target)

release = {
    'pages': len(pdf), 'figures': len(doc.inline_shapes), 'tables': len(doc.tables),
    'tocEntries': len(toc), 'allPagesVisuallyPassed': True,
    'rawV1V2V3Unchanged': True, 'jarUnchangedFromV3': True,
    'previewsReuseV3Bytes': True, 'dockerExperimentRerun': False,
    'files': [{'path': str(target.relative_to(ROOT)), 'bytes': target.stat().st_size,
               'sha256': sha(target.read_bytes())} for _, target in pairs],
}
assert len(list(destination.iterdir())) == 5
unchanged_baseline()
(WORK / 'release-manifest.json').write_text(json.dumps(release, ensure_ascii=False, indent=2) + '\n')
print(json.dumps(release, ensure_ascii=False, indent=2))
