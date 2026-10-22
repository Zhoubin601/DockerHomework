from pathlib import Path
from copy import deepcopy
from zipfile import ZipFile
import hashlib, json, re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image

WORK=Path(__file__).resolve().parent
ROOT=WORK.parents[1]
SOURCE=ROOT/'output/v6/memory-game6085.docx'
DEST=WORK/'20261022-docker-lab-report-v7.docx'

STAGES=[
 ('01-before-install','安装前检查','在新建的 Ubuntu 环境中检查系统版本、CPU 架构、Docker 命令和 socket。检查时 docker、dockerd 尚未安装，/var/run/docker.sock 不存在。'),
 ('02-official-repository','配置 Docker 官方软件源','先安装 ca-certificates 和 curl，再保存 Docker 官方 GPG 公钥，添加 noble stable 软件源并更新软件包索引。apt-cache 显示此时 Docker CE 尚未安装，候选版本为 29.9.2。'),
 ('03-install-engine','安装 Docker Engine 和插件','通过 apt-get 安装 docker-ce、docker-ce-cli、containerd.io、Buildx 和 Compose 插件。软件包配置完成后，分别查看包版本及 docker、dockerd、Compose 的版本，实际安装 Docker Engine 29.9.2 和 Compose 5.6.1。'),
 ('04-first-configuration','首次配置与引擎启动','创建 /etc/docker/daemon.json，设置独立数据目录、vfs 存储驱动及 json-file 日志轮转。vfs 用于本次嵌套实验环境。dockerd --validate 返回 configuration OK 后，启动新引擎并核对 Client、Server、数据目录与日志驱动。'),
 ('05-verify-installation','导入基础镜像并验证安装','为复用已下载的基础镜像，先在外层导出 Temurin、Registry 和 hello-world，再通过 docker load 导入本轮独立引擎。运行 hello-world 后出现 Hello from Docker，容器以退出码 0 结束，验证新安装的引擎能够实际运行容器。'),
 ('06-create-resources','创建网络 数据卷和镜像仓库','依次创建 memory-game6085-v7-net 网络、应用数据卷和仓库数据卷，再启动 Registry 容器。inspect 核对 bridge 网络及 local 数据卷驱动，仓库端口绑定补充环境内部的 127.0.0.1:5000，/v2/ 接口返回空 JSON 对象。'),
 ('07-build-application','构建应用镜像并核对 JAR','采用源码包内的 Dockerfile.prebuilt 和配套 JAR，在新安装的引擎内构建 memory-game6085:1.0.2-v7。该步骤复用已打包的应用，不重新执行正文的 npm 与 Maven 构建。随后从容器复制 /app/app.jar，与配套 JAR 计算 SHA-256 并逐字节比较，结果一致。'),
 ('08-run-application','配置应用容器并保存成绩','启动应用容器时明确指定网络、18085:8080 端口映射、DATA_DIR 和数据卷挂载。健康接口返回 UP，运行用户为 garden，UID/GID 为 999。外层把本机 127.0.0.1:18089 转发到该环境的 18085，浏览器通过此地址完成一局轻松模式并保存成绩。'),
 ('09-publish-image','标记并发布镜像','保存原容器标识和非空排行榜后，将镜像标记为 localhost:5000/memory-game6085:1.0.2-v7，并推送到本轮仓库。push 输出 Pushed 和镜像摘要，仓库标签接口返回 1.0.2-v7。'),
 ('10-remove-and-pull','删除本地标签后重新拉取','移除本轮应用容器和 JAR 校验容器，再删除应用镜像的两个本地标签。inspect 确认标签不存在、镜像列表为空后，才从 Registry 执行 pull。拉取返回 Downloaded newer image，目标镜像摘要与 push 一致。构建缓存及其他基础镜像未作为本地应用标签存在的依据。'),
 ('11-recreate-and-verify','重建容器并核对持久化数据','使用拉取的仓库镜像创建 garden6085-pulled，重新挂载同一应用数据卷。新旧容器标识不同，完整排行榜 JSON 逐字节一致，输出 FullLeaderboardPreserved=True。最后在新容器的浏览器页面再次查看原成绩。'),
]

def sha(data): return hashlib.sha256(data).hexdigest()
results={}
for slug,_,_ in STAGES:
 result=json.loads((WORK/(slug+'-result.json')).read_text(encoding='utf-8'))
 assert result['exit_code']==0,(slug,result)
 results[slug]=result
for mode in ['before','after']:
 assert json.loads((WORK/f'browser-{mode}.json').read_text())['result']=='passed'
score=json.loads((WORK/'saved-score.json').read_text(encoding='utf-8'))
assert score['nickname']=='6085补充实验'
before=(WORK/'before-leaderboard.json').read_bytes()
assert before==(WORK/'after-leaderboard.json').read_bytes()
assert json.loads(before)['scores'][0]==score
before_id=(WORK/'before-container.txt').read_text().strip()
after_id=(WORK/'after-container.txt').read_text().strip()
assert before_id!=after_id
push=re.search(r'digest:\s*(sha256:[0-9a-f]{64})',(WORK/'09-push-full.log').read_text()).group(1)
pull=re.search(r'Digest:\s*(sha256:[0-9a-f]{64})',(WORK/'10-pull-full.log').read_text()).group(1)
assert push==pull

doc=Document(SOURCE)
changes=[]
def replace_text(p,new):
 nodes=p._p.xpath('.//w:t')
 old=''.join(x.text or '' for x in nodes)
 assert nodes
 nodes[0].text=new
 for n in nodes[1:]:n.text=''
 changes.append({'before':old,'after':new})

cover=next(p for p in doc.paragraphs if p.text=='实验日期：2026年10月21日')
replace_text(cover,'实验日期：2026年10月21日；补充实验：10月22日')
for p in doc.paragraphs:
 if p.text.startswith('第二至九章记录原 Windows 实验。'):
  replace_text(p,'第二至九章记录 2026 年 10 月 21 日的 Windows 实验。附录 A 记录 10 月 22 日在本机隔离 Ubuntu 环境中新安装 Docker Engine 及后续部署的补充实验，环境、资源和结果分别列示。')
 elif p.text.startswith('原 Windows 环境沿用既有安装。本报告附录A'):
  replace_text(p,'原 Windows 环境沿用既有安装，历史安装过程截图未提供。附录 A 补充 2026 年 10 月 22 日在本机隔离 Ubuntu 环境中实际安装 Docker Engine、首次配置以及网络、数据卷和容器创建的操作记录。')
 elif p.text.startswith('项目源代码：Code/memory-game6085/；实验结果：'):
  replace_text(p,'源码包解压后的项目目录为 memory-game6085/。正文图 1 至图 11 来自原 Windows 实验，其结果记录在 docs/docker-evidence/results.json；附录 A 来自本轮补充实验，命令、完整日志和原图记录在 docs/v7-work/。报告中的过程截图均为内嵌图片。')
 elif p.text=='技术资料查阅日期：2026年10月21日。':
  replace_text(p,'正文技术资料查阅日期为 2026 年 10 月 21 日；Docker Engine 安装文档核对日期为 10 月 22 日。')

mapping={
 'Spring Boot 项目打包发布':'Windows 多阶段构建见第五章；补充环境的配套 JAR 镜像构建见附录 A7。',
 '发布后演示拉取':'Windows 仓库发布与拉取见第六章；补充环境的发布与拉取见附录 A9 至 A10。',
 '全过程截图汇总':'正文图1至图11，以及附录图A1至图A11（含分图），均嵌入本Word。',
 'Docker 环境部署与验证':'Windows 版本验证见2.2；补充环境的安装、首次配置、验证及资源创建见附录A1至A6。',
}
for row in doc.tables[0].rows:
 if row.cells[0].text in mapping:replace_text(row.cells[1].paragraphs[0],mapping[row.cells[0].text])

doc.add_paragraph('Docker Engine 官方 Ubuntu 安装文档：https://docs.docker.com/engine/install/ubuntu/',style='Reference')

def body(text):
 p=doc.add_paragraph(text)
 p.paragraph_format.line_spacing=1.15
 p.paragraph_format.space_after=Pt(6)
 for r in p.runs:r.font.size=Pt(11)
 return p

def code(text):
 p=doc.add_paragraph(text,style='Code')
 p.paragraph_format.first_line_indent=Pt(0)
 p.paragraph_format.space_after=Pt(5)
 return p

figure_records=[]
def picture(file,caption):
 p=doc.add_paragraph()
 p.paragraph_format.first_line_indent=Pt(0)
 p.paragraph_format.space_after=Pt(3)
 p.paragraph_format.keep_with_next=True
 p.alignment=WD_ALIGN_PARAGRAPH.CENTER
 with Image.open(file) as im:w,h=im.size
 width=min(6.2,6.4*w/h)
 run=p.add_run()
 shape=run.add_picture(str(file),width=Inches(width))
 shape._inline.docPr.set('descr',caption)
 cp=doc.add_paragraph(caption,style='题注1')
 cp.alignment=WD_ALIGN_PARAGRAPH.CENTER
 cp.paragraph_format.first_line_indent=Pt(0)
 cp.paragraph_format.space_after=Pt(8)
 cp.paragraph_format.keep_together=True
 figure_records.append({'caption':caption,'source':str(file.relative_to(ROOT)),'sha256':sha(file.read_bytes())})

h=doc.add_paragraph('附录 A Docker 安装与配置补充实验',style='Heading 1')
h.paragraph_format.page_break_before=True
h.paragraph_format.first_line_indent=Pt(0)
for r in h.runs:r.font.color.rgb=RGBColor(0,0,0)
body('本附录记录 2026 年 10 月 22 日在本机独立完成的 Linux 补充实验，从未安装 Docker 的 Ubuntu 环境开始，依次完成软件源配置、Docker Engine 安装、首次配置、资源创建和应用发布拉取。原 Windows 实验与本轮补充实验使用不同引擎和资源。')
body('外层环境是本机 Docker Desktop，提供隔离 Ubuntu 容器；Docker Engine 29.9.2 安装在该容器内部，未挂载宿主机 Docker socket。后续应用镜像、网络、数据卷和 Registry 均由新安装的内层引擎管理。该环境用于课程补充实验。')
body('图中的安装和命令操作记录为实际执行日志在浏览器中的截屏，按原日志顺序保留命令、输出与退出码；游戏页面图直接采集自本轮运行的应用容器。日志中的时间使用 UTC，与北京时间相差 8 小时。')

table=doc.add_table(rows=1,cols=2)
table.autofit=False
table.columns[0].width=Inches(1.45);table.columns[1].width=Inches(4.75)
table.rows[0].cells[0].text='项目';table.rows[0].cells[1].text='本轮补充实验实际配置'
summary=[
 ('安装系统','Ubuntu 24.04.5 LTS，amd64；全新隔离环境内安装'),
 ('Docker 组件','Engine / CLI 29.9.2；Compose 5.6.1\nBuildx 0.38.1；containerd.io 2.4.2'),
 ('引擎配置','数据目录 /var/lib/docker-lab6085；vfs；json-file 日志轮转'),
 ('应用与仓库','镜像 memory-game6085:1.0.2-v7；内层 Registry localhost:5000'),
 ('网络和数据卷','memory-game6085-v7-net\nmemory-game6085-v7-data\nmemory-game6085-v7-registry'),
 ('访问路径','本机127.0.0.1:18089 → 隔离环境18085 → 应用容器8080'),
 ('应用JAR','沿用源码包配套JAR；容器内外SHA-256及字节比较一致'),
 ('演示成绩',f'{score["nickname"]}；编号{score["id"]}；{score["moves"]}步；{score["durationMs"]} ms'),
 ('重建验证',f'{before_id[:12]} → {after_id[:12]}；完整排行榜一致'),
]
for left,right in summary:
 row=table.add_row();row.cells[0].text=left;row.cells[1].text=right
for i,row in enumerate(table.rows):
 row.cells[0].width=Inches(1.45)
 row.cells[1].width=Inches(4.75)
 for cell in row.cells:
  for p in cell.paragraphs:
   p.paragraph_format.first_line_indent=Pt(0)
   p.paragraph_format.space_after=Pt(4)
   p.paragraph_format.space_before=Pt(4)
   for r in p.runs:r.font.size=Pt(10)
  if i==0:
   shading=OxmlElement('w:shd');shading.set(qn('w:val'),'clear');shading.set(qn('w:fill'),'E9ECEF');cell._tc.get_or_add_tcPr().append(shading)
   for r in cell.paragraphs[0].runs:r.bold=True
 for old in row._tr.xpath('./w:trPr/w:cantSplit'):old.getparent().remove(old)
 row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))

for index,(slug,title,description) in enumerate(STAGES,1):
 h=doc.add_paragraph(f'A{index} {title}',style='Heading 2')
 h.paragraph_format.page_break_before=True
 h.paragraph_format.first_line_indent=Pt(0)
 for r in h.runs:r.font.color.rgb=RGBColor(0,0,0)
 body(description)
 if index==7:
  code('docker build --pull=false --network=none -f Dockerfile.prebuilt -t memory-game6085:1.0.2-v7 .')
 if index==8:
  body(f'本次保存的记录编号为 {score["id"]}，昵称为“{score["nickname"]}”，成绩为 {score["moves"]} 步、{score["durationMs"]} ms。启动健康检查曾早于应用就绪，补充重试后通过；该次初始失败记录另行保存在过程文档中。')
 if index==11:
  code(push)
  body(f'原容器标识前12位为 {before_id[:12]}，新容器为 {after_id[:12]}。重建前后保留同一条成绩，昵称、难度、步数、用时和创建时间均一致。')
 captures=json.loads((WORK/(slug+'-captures.json')).read_text())
 extras=[]
 if index==8:extras=[('08-browser-saved.png','浏览器完成游戏并保存成绩')]
 if index==11:extras=[('11-browser-recreated.png','拉取镜像重建后的浏览器排行榜')]
 count=len(captures)+len(extras)
 for j,name in enumerate(captures):
  label=f'图A{index}'+(chr(97+j) if count>1 else '')
  picture(WORK/name,f'{label} {title}'+(f' 操作日志第{j+1}部分' if len(captures)>1 else ' 操作日志'))
 for j,(name,description2) in enumerate(extras,len(captures)):
  picture(WORK/name,f'图A{index}{chr(97+j)} {description2}')

body('本轮从空白 Linux 环境安装 Docker Engine 后完成应用镜像构建、发布和拉取。新容器标识发生变化，原数据卷及已保存成绩保留；完整排行榜比较和浏览器结果相互对应。正文的原 Windows 实验记录与本轮补充记录分别保留。')
doc.core_properties.revision=7
doc.core_properties.description='Spring Boot 小游戏 Docker 容器化实验报告 含独立Linux安装配置与发布拉取补充实验'
for n in doc.element.xpath('.//w:lastRenderedPageBreak'):n.getparent().remove(n)
doc.save(DEST)

with ZipFile(SOURCE) as a,ZipFile(DEST) as b:
 old=[sha(a.read(n)) for n in a.namelist() if n.startswith('word/media/')]
 new=[sha(b.read(n)) for n in b.namelist() if n.startswith('word/media/')]
 assert len(old)==11 and all(x in new for x in old)
assert len(doc.inline_shapes)==11+len(figure_records)
text='\n'.join(p.text for p in doc.paragraphs)+'\n'+'\n'.join(c.text for t in doc.tables for r in t.rows for c in r.cells)
assert '图B1' not in text and '附录A5—A8' not in text and 'localhost:15085' not in text
assert len([p for p in doc.paragraphs if p.text.startswith('附录 A') and p.style.name=='Heading 1'])==1
(WORK/'report-build.json').write_text(json.dumps({'source':str(SOURCE),'candidate':str(DEST),'source_sha256':sha(SOURCE.read_bytes()),'images':len(doc.inline_shapes),'tables':len(doc.tables),'text_changes':changes,'new_figures':figure_records},ensure_ascii=False,indent=2),encoding='utf-8')
(WORK/'results.json').write_text(json.dumps({'stages':results,'engine':'29.9.2','compose':'5.6.1','image':'localhost:5000/memory-game6085:1.0.2-v7','push_digest':push,'pull_digest':pull,'before_container':before_id,'after_container':after_id,'score':score,'full_leaderboard_preserved':True,'base':'http://127.0.0.1:18089'},ensure_ascii=False,indent=2),encoding='utf-8')
print(f'Candidate created: {DEST.name}; images={len(doc.inline_shapes)}; tables={len(doc.tables)}')
