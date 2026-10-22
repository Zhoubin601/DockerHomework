from pathlib import Path
from zipfile import ZipFile
import importlib.util, json, hashlib, re
import pdfplumber
from docx import Document
from PIL import Image, ImageDraw

work=Path(__file__).resolve().parent
root=work.parents[1]
helper=Path('C:/Users/a3185/.codex/plugins/cache/openai-primary-runtime/documents/26.1020.11127/skills/documents/render_docx.py')
spec=importlib.util.spec_from_file_location('skill_render_docx',helper)
module=importlib.util.module_from_spec(spec);spec.loader.exec_module(module)
# Windows runtime has no LibreOffice. Use the checked native Word PDF for the
# skill's unchanged Poppler rasterization stage; do not edit the managed helper.
module.convert_to_pdf=lambda *args,**kwargs:(str(work/'native-final.pdf'),'Native Word COM export; Windows LibreOffice unavailable')
dest=work/'render-final'
rendered=module.rasterize(str(work/'word-final.docx'),str(dest),144,False,False)

native=json.loads((work/'word-native-check.json').read_text(encoding='utf-8-sig'))
build=json.loads((work/'report-build.json').read_text(encoding='utf-8'))
norm=lambda s:re.sub(r'\s+','',s)
document=Document(work/'word-final.docx')
paragraphs=[p.text for p in document.paragraphs]
all_text='\n'.join(paragraphs)+'\n'+'\n'.join(c.text for t in document.tables for r in t.rows for c in r.cells)
captions=[p for p in paragraphs if re.match(r'^图(?:\d+|A\d+)',p)]
with pdfplumber.open(work/'native-final.pdf') as pdf:
 texts=[p.extract_text() or '' for p in pdf.pages]
 normalized=[norm(t) for t in texts]
 pages=[]
 for i,page in enumerate(pdf.pages):
  words=page.extract_words()
  clipped=[w['text'] for w in words if w['x0'] < -1 or w['x1'] > page.width+1 or w['top'] < -1 or w['bottom'] > page.height+1]
  pages.append({'page':i+1,'text_characters':len(texts[i]),'images':len(page.images),'captions':[c for c in captions if norm(c) in normalized[i]],'words_outside_page':clipped})
 checks=[{**h,'present_on_native_page':norm(h['heading']) in normalized[h['page']-1]} for h in native['headings']]
 assert all(x['present_on_native_page'] for x in checks)
 assert len(pdf.pages)==native['pages']==len(rendered)
 assert all(x['text_characters']>25 and not x['words_outside_page'] for x in pages)
 assert all(any(norm(c) in t for t in normalized) for c in captions)
 assert all(any(norm(c) in normalized[i] and len(pdf.pages[i].images)>0 for i in range(len(pdf.pages))) for c in captions)

toc_lines=[s for toc in native['tocText'] for s in toc.split('\r') if s.strip()]
toc=[]
for line in toc_lines:
 parts=line.strip().split('\t')
 if len(parts)>1 and parts[-1].strip().isdigit():
  toc.append({'heading':parts[0].strip(),'page':int(parts[-1])})
assert len(toc)==len(checks),(len(toc),len(checks),toc_lines)
for item in toc:
 assert any(norm(item['heading'])==norm(h['heading']) and item['page']==h['page'] for h in checks),item

sha=lambda b:hashlib.sha256(b).hexdigest()
with ZipFile(root/'output/v6/memory-game6085.docx') as z:
 old=[sha(z.read(n)) for n in z.namelist() if n.startswith('word/media/')]
with ZipFile(work/'word-final.docx') as z:
 media={n:sha(z.read(n)) for n in z.namelist() if n.startswith('word/media/')}
 assert z.testzip() is None
 assert all(x in media.values() for x in old)
 assert all(x['sha256'] in media.values() for x in build['new_figures'])
 assert len(z.namelist())>0
assert len(document.inline_shapes)==31 and len(captions)==31
assert '图B1' not in all_text and 'localhost:15085' not in all_text
assert not any(x in all_text for x in ['待确认','TODO','占位符','附录A2—A3','附录A5—A8'])

for start in range(1,len(rendered)+1,9):
 end=min(start+8,len(rendered))
 sheet=Image.new('RGB',(1380,1935),'#d8d8d8');draw=ImageDraw.Draw(sheet)
 for offset,index in enumerate(range(start,end+1)):
  im=Image.open(dest/f'page-{index}.png');im.thumbnail((440,605))
  x=offset%3*460+10;y=offset//3*645+25
  sheet.paste(im,(x,y));draw.text((x,y-18),f'Page {index}',fill='black')
 sheet.save(dest/f'contact-{start}-{end}.jpg',quality=90)

result={'pages':native['pages'],'headings_and_toc':len(checks),'images':len(document.inline_shapes),'tables':len(document.tables),'captions':len(captions),'old_images_preserved':len(old),'new_images_matched':len(build['new_figures']),'heading_pages_verified':checks,'toc_entries_verified':toc,'page_checks':pages,'media_sha256':media,'final_sha256':sha((work/'word-final.docx').read_bytes()),'renderer':'documents render_docx.py Poppler stage with native Word PDF adapter; bundled Windows LibreOffice unavailable'}
(work/'final-checks.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
(work/'final-text.txt').write_text(all_text,encoding='utf-8')
print(json.dumps({k:result[k] for k in ['pages','headings_and_toc','images','tables','captions','old_images_preserved','new_images_matched']},ensure_ascii=False))
