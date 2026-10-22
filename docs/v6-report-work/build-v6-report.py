"""Apply the approved language edits to a snapshot of the user's current v5."""
from pathlib import Path
from zipfile import ZipFile
import hashlib
import json

from docx import Document
from docx.oxml.ns import qn

WORK = Path(__file__).resolve().parent
SOURCE = WORK / 'source-v5.docx'
DEST = WORK / '20261022-docker-lab-report-v6.docx'
doc = Document(SOURCE)
changes = []
by_id = {p._p.get(qn('w14:paraId')): p for p in doc.paragraphs}


def text_of(p):
    return ''.join(n.text or '' for n in p._p.xpath('.//w:t'))


def replace_in_paragraph(p, old, new, group):
    """Keep paragraph properties, bookmarks and fields; edit visible text only."""
    nodes = p._p.xpath('.//w:t')
    before = ''.join(n.text or '' for n in nodes)
    assert before.count(old) == 1, (group, old, before)
    start = before.index(old)
    end = start + len(old)
    offset = 0
    inserted = False
    for node in nodes:
        value = node.text or ''
        lo, hi = offset, offset + len(value)
        offset = hi
        if hi <= start or lo >= end:
            continue
        left = value[:max(0, start - lo)]
        right = value[max(0, end - lo):] if end < hi else ''
        node.text = left + (new if not inserted else '') + right
        inserted = True
        if node.text.startswith(' ') or node.text.endswith(' '):
            node.set(qn('xml:space'), 'preserve')
    assert inserted
    after = text_of(p)
    assert after == before.replace(old, new, 1), group
    changes.append({'group': group, 'paraId': p._p.get(qn('w14:paraId')),
                    'before': before, 'after': after})


def replace_pid(pid, old, new, group):
    replace_in_paragraph(by_id[pid], old, new, group)


def replace_cell(table_index, row_label, column, old, new, group):
    table = doc.tables[table_index - 1]
    row = next(r for r in table.rows if r.cells[0].text == row_label)
    cell = row.cells[column - 1]
    assert cell.text == old
    assert len(cell.paragraphs) == 1
    replace_in_paragraph(cell.paragraphs[0], old, new, group)


replace_pid('63C858E2', 'v5规范文档结构及排版，三轮实际版本、端口和结果分别保留。', '', 1)
replace_pid('6F2333DD', text_of(by_id['6F2333DD']),
            '第二至九章记录原 Windows 实验。Linux 补充部署使用独立资源，环境配置和操作截图见附录 A。', 2)
old_title = '4.2 容器中的真实游戏操作'
new_title = '4.2 容器中的游戏操作与成绩保存'
replace_pid('7B2E61BC', old_title, new_title, 3)
# Update the existing cached TOC title as well; Word will recalculate its fields.
replace_pid('688B7EBD', old_title, new_title, 3)
replace_pid('3D686D3F', text_of(by_id['3D686D3F']),
            '通过浏览器在应用容器中完成轻松难度的一局游戏，并以昵称“6085实验”保存成绩。服务端返回的记录编号为 1，成绩为 12 步、6.95 秒，用于演示成绩保存流程。', 3)
replace_pid('11ACE681', '程序对重建前后的完整排行榜 JSON 比较，结果为 LeaderboardPreserved=True。',
            '程序比较了重建前后的完整排行榜 JSON，结果为 LeaderboardPreserved=True。', 4)
replace_pid('1911E03C', '恢复后以 Client 与 Server 响应以及实际容器结果验证环境可用。',
            '恢复后，Client 与 Server 均能响应，容器也可正常运行。', 5)
replace_pid('45CBD8A4', '重建前后容器标识发生变化，但保存的成绩没有变化，说明 SQLite 文件与应用容器已经通过命名数据卷实现生命周期分离。',
            '重建前后，容器标识发生变化，已保存的成绩保持不变，说明容器重建后仍能读取原数据卷中的 SQLite 数据。', 6)
replace_cell(6, 'SQLite 数据卷', 2, '完整榜单对比一致，保存成绩保留',
             '重建前后的完整榜单一致，已保存成绩未丢失', 7)
replace_pid('6270524A', text_of(by_id['6270524A']),
            '分别计算配套 JAR 和容器内 /app/app.jar 的 SHA256，结果一致，确认容器内的 JAR 与配套 JAR 相同。', 8)
replace_pid('3E298FFB', 'v5独立桌面及手机预览采用10月22日本机复测截图，两轮采集来源分别记录。', '', 9)
replace_pid('53698D40', text_of(by_id['53698D40']),
            '正文图 1 至图 11 来自原 Windows 实验；本附录图 A1a、图 A1b 及图 A2 至图 A9 来自 2026 年 10 月 21 日的 Linux 补充实验。终端截图采集自当次窗口，浏览器截图采集自当次应用容器。', 10)
replace_pid('7E6FB9DB', 'Windows原实验镜像为1.0.0，Linux补充镜像为1.0.1-v3；v5修订报告结构，应用JAR及这些历史实验结果不变。',
            'Windows 原实验镜像版本为 1.0.0，Linux 补充实验镜像版本为 1.0.1-v3。', 11)
replace_cell(8, 'A3', 2, '运行与JAR一致', '容器运行与 JAR 一致性校验', 12)
replace_pid('5CD5D184', text_of(by_id['5CD5D184']),
            '本轮源码镜像 JAR 与配套 JAR 的整包摘要不同，但 BOOT-INF 内的应用类、前端资源和依赖逐项相同。完整命令、本地镜像移除检查及输出收录在证据包 windows-localtest-20261022 目录中。', 13)
replace_cell(9, '浏览器与用户', 2,
             '实际容器Playwright 3/3；应用以garden普通用户运行，健康检查UP。',
             '容器页面的 Playwright 测试 3/3 通过；应用以普通用户 garden 运行，健康检查返回 UP。', 14)
replace_cell(9, '真实测试成绩', 1, '真实测试成绩', '测试成绩', 14)
replace_pid('35D6267B', text_of(by_id['35D6267B']),
            '证据来源：docs/local-test-20261022/docker-roundtrip/results.json 及 after-recreate-ui.png；构建与浏览器测试日志同收录于证据包。本轮沿用已安装的 Docker，未重新安装。', 15)

# Keep the user's latest cover date; do not restore the old retest suffix.
assert text_of(by_id['52BDA62F']) == '实验日期：2026年10月21日'
assert set(x['group'] for x in changes) == set(range(1, 16))
for node in doc.element.xpath('.//w:lastRenderedPageBreak'):
    node.getparent().remove(node)
doc.core_properties.description = 'Spring Boot 小游戏 Docker 容器化实验报告'
doc.core_properties.revision = 6
doc.save(DEST)

with ZipFile(SOURCE) as a, ZipFile(DEST) as b:
    media_a = {n: hashlib.sha256(a.read(n)).hexdigest() for n in a.namelist() if n.startswith('word/media/')}
    media_b = {n: hashlib.sha256(b.read(n)).hexdigest() for n in b.namelist() if n.startswith('word/media/')}
    assert media_a == media_b and len(media_a) == 22
assert len(doc.tables) == 9 and len(doc.inline_shapes) == 22
log = {'source': str(SOURCE), 'source_sha256': hashlib.sha256(SOURCE.read_bytes()).hexdigest(),
       'candidate': str(DEST), 'groups_applied': sorted({x['group'] for x in changes}),
       'text_changes': changes, 'media_unchanged': True,
       'figures': len(doc.inline_shapes), 'tables': len(doc.tables),
       'cover_date_preserved': text_of(by_id['52BDA62F'])}
(WORK / 'applied-changes.json').write_text(json.dumps(log, ensure_ascii=False, indent=2), encoding='utf-8')
print(f'Applied 15 groups / {len(changes)} text edits; 22 images and 9 tables retained.')
