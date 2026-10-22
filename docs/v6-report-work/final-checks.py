from collections import Counter
from pathlib import Path
from zipfile import ZipFile
from urllib.parse import unquote
import hashlib
import json
import posixpath
import re

from lxml import etree as E
from docx import Document
from docx.oxml.ns import qn

work = Path(__file__).resolve().parent
root = work.parents[1]
source = work / 'source-v5.docx'
candidate = work / '20261022-docker-lab-report-v6.docx'
ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main',
      'a': 'http://schemas.openxmlformats.org/drawingml/2006/main'}

sha = lambda p: hashlib.sha256(p.read_bytes()).hexdigest()
data = {'candidate_sha256': sha(candidate), 'source_sha256': sha(source)}
with ZipFile(source) as a, ZipFile(candidate) as b:
    media = lambda z: Counter(hashlib.sha256(z.read(n)).hexdigest()
                              for n in z.namelist() if n.startswith('word/media/'))
    data['all_22_images_unchanged'] = media(a) == media(b) and sum(media(b).values()) == 22
    data['zip_crc_ok'] = b.testzip() is None
    errors = []
    names = set(b.namelist())
    for name in names:
        if not name.endswith(('.xml', '.rels')):
            continue
        part = E.fromstring(b.read(name))
        if name.endswith('.rels'):
            base = posixpath.dirname(posixpath.dirname(name)) if '/_rels/' in name else ''
            for rel in part:
                if rel.get('TargetMode') == 'External':
                    continue
                target = unquote(rel.get('Target', '').split('#', 1)[0])
                resolved = target.lstrip('/') if target.startswith('/') else posixpath.normpath(posixpath.join(base, target))
                if resolved and resolved not in names:
                    errors.append({'part': name, 'target': resolved})
    data['missing_relationship_targets'] = errors
    tree = E.fromstring(b.read('word/document.xml'))
    alltext = '\n'.join(''.join(p.xpath('.//w:t/text()', namespaces=ns)) for p in tree.xpath('.//w:p', namespaces=ns))
    data['unresolved_markers'] = [s for s in ['待确认', 'TODO', 'TBD', '{{', 'lorem ipsum', '未补造', '未把整包摘要与应用内容校验混同', 'v5', 'v6'] if s in alltext]
    data['revision_count'] = len(tree.xpath('.//w:ins|.//w:del|.//w:moveFrom|.//w:moveTo', namespaces=ns))
    fields = tree.xpath('.//w:instrText/text()', namespaces=ns)
    data['live_toc_exists'] = any('TOC ' in s for s in fields)
    footer_fields = []
    for name in names:
        if name.startswith('word/footer') and name.endswith('.xml'):
            footer_fields += E.fromstring(b.read(name)).xpath('.//w:instrText/text()', namespaces=ns)
    data['live_page_field_exists'] = any(re.search(r'\bPAGE\b', s) for s in footer_fields)

doc = Document(candidate)
source_doc = Document(source)
data['tables'] = len(doc.tables)
data['figures'] = len(doc.inline_shapes)
data['table_shapes_unchanged'] = [(len(t.rows), len(t.columns)) for t in doc.tables] == [(len(t.rows), len(t.columns)) for t in source_doc.tables]
data['cover_date_preserved'] = next(p.text for p in doc.paragraphs if p.text.startswith('实验日期：')) == '实验日期：2026年10月21日'
data['native_pagination'] = json.loads((work / 'native-pagination-verification.json').read_text(encoding='utf-8'))
data['strict_schema_passed'] = 'no errors found' in (work / 'schema-final.txt').read_text(encoding='utf-8-sig')

# Resolve every first-line indent warning through the existing style hierarchy.
issue_log = (work / 'issues-final.txt').read_text(encoding='utf-8-sig')
format_ids = re.findall(r'\[F\d+\] /body/p\[@paraId=([^\]]+)\]: Body paragraph missing first-line indent', issue_log)
code_ids = re.findall(r'\[C\d+\] /body/p\[@paraId=([^\]]+)\]/r\[\d+\]: Consecutive spaces', issue_log)
by_id = {p._p.get(qn('w14:paraId')): p for p in doc.paragraphs}
indent_checks = []
for pid in format_ids:
    p = by_id[pid]
    indent = p.paragraph_format.first_line_indent
    style = p.style
    while indent is None and style is not None:
        indent = style.paragraph_format.first_line_indent
        style = style.base_style
    indent_checks.append({'paraId': pid, 'effective_twips': indent.twips if indent is not None else None})
data['format_warning_checks'] = indent_checks
data['all_code_spacing_warnings_intentional'] = all(by_id[pid].style.name == 'Code' for pid in code_ids)
data['format_warning_count'] = len(format_ids)
data['code_spacing_warning_count'] = len(code_ids)

baseline = json.loads((work / 'baseline-sha256.json').read_text(encoding='utf-8-sig'))
data['historical_baseline_files'] = len(baseline)
data['changed_historical_files'] = [x['Path'] for x in baseline if sha(Path(x['Path'])) != x['SHA256'].lower()]
data['visual_review_files'] = [f'visual-qa-pages{span}.md' for span in ['1-10', '11-20', '21-31']]
data['all_visual_reviews_present'] = all((work / p).exists() for p in data['visual_review_files'])
data['passed'] = all([
    data['all_22_images_unchanged'], data['zip_crc_ok'], not errors, not data['unresolved_markers'],
    data['revision_count'] == 0, data['live_toc_exists'], data['live_page_field_exists'],
    data['tables'] == 9, data['figures'] == 22, data['table_shapes_unchanged'], data['cover_date_preserved'],
    data['native_pagination']['tocEntries'] == 42, data['native_pagination']['allTOCPagesVerifiedInPDF'],
    data['strict_schema_passed'], all((p['effective_twips'] or 0) > 0 for p in indent_checks),
    len(format_ids) == 31, len(code_ids) == 14, data['all_code_spacing_warnings_intentional'],
    not data['changed_historical_files'], data['all_visual_reviews_present']])
(work / 'final-checks.json').write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding='utf-8')
print(json.dumps({k: data[k] for k in ['candidate_sha256', 'tables', 'figures', 'strict_schema_passed',
      'format_warning_count', 'code_spacing_warning_count', 'historical_baseline_files', 'changed_historical_files', 'all_visual_reviews_present', 'passed']}, indent=2))
raise SystemExit(0 if data['passed'] else 1)
