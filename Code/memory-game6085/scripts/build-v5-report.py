"""Repair v4 OOXML, preserve original evidence, and add the verified local retest."""
from pathlib import Path
from copy import deepcopy
from collections import Counter
from zipfile import ZipFile, ZIP_DEFLATED
import ast
import inspect
import json
import re

from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.oxml.text.parfmt import CT_PPr
from docx.oxml.text.font import CT_RPr
from docx.oxml.table import CT_TcPr
from docx.oxml.settings import CT_Settings
from lxml import etree as E

ROOT=Path(__file__).resolve().parents[3]
WORK=ROOT/'docs/v5-report-work'
WORK.mkdir(exist_ok=True)
SOURCE=ROOT/'output/v4/20261021-docker-lab-report-v4.docx'
DEST=WORK/'20261022-docker-lab-report-v5.docx'
LOCAL=ROOT/'docs/local-test-20261022'
RESULT=json.loads((LOCAL/'docker-roundtrip/results.json').read_text(encoding='utf-8'))
assert all(RESULT[k] for k in ['containerChanged','leaderboardPreserved','repoDigestsMatch','originalContainersUnchanged'])
doc=Document(SOURCE)

def set_text(p,text):
    formatting=deepcopy(p.runs[0]._r.rPr) if p.runs and p.runs[0]._r.rPr is not None else None
    p.clear()
    run=p.add_run(text)
    if formatting is not None:run._r.insert(0,formatting)

def fonts(run,ea='宋体',latin='Times New Roman',size=11,bold=False):
    run.font.name=latin
    run.font.size=Pt(size)
    run.font.bold=bold
    run.font.color.rgb=RGBColor(0,0,0)
    run._r.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),ea)

for name,base in [('Caption','Normal'),('Code','Normal'),('Reference','Normal'),('toc 1','Normal'),('toc 2','Normal'),('Heading 3','Normal')]:
    if name not in doc.styles:
        style=doc.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH)
        style.base_style=doc.styles[base]

for name in ['Normal','Caption','Code','Reference','toc 1','toc 2']:
    style=doc.styles[name]
    style.font.name='Times New Roman' if name!='Code' else 'Consolas'
    style.font.size=Pt(11 if name=='Normal' else (8.5 if name=='Code' else 10))
    style.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'宋体')
    style.paragraph_format.first_line_indent=Cm(.74 if name=='Normal' else 0)
for name,size in [('Heading 1',18),('Heading 2',14),('Heading 3',12)]:
    doc.styles[name].font.size=Pt(size)
    doc.styles[name].font.bold=True
    doc.styles[name].font.name='黑体'
    doc.styles[name].element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'黑体')

for n in doc.element.xpath('.//w:lastRenderedPageBreak'):
    n.getparent().remove(n)
old_ids=[]
for n in doc.element.xpath('.//w:bookmarkStart'):
    if re.match(r'V[234]Heading',n.get(qn('w:name'),'')):
        old_ids.append(n.get(qn('w:id')))
        n.getparent().remove(n)
for n in doc.element.xpath('.//w:bookmarkEnd'):
    if n.get(qn('w:id')) in old_ids:n.getparent().remove(n)

for p in doc.paragraphs:
    if p.text.startswith('本实验将 Vue3'):
        set_text(p,'本实验将 Vue3、Spring Boot 与 SQLite 记忆翻牌小游戏打包为 Docker 镜像，完成本机仓库发布、删除本地镜像后拉取和数据卷保留验证。正文记录2026年10月21日 Windows 实验，附录A记录当日 Linux 补充实验，附录B记录2026年10月22日本机复测。v5规范文档结构及排版，三轮实际版本、端口和结果分别保留。')
    elif p.text=='实验日期：2026年10月21日':
        set_text(p,'实验日期：2026年10月21日；复测：10月22日')
    elif p.text.startswith('正文图1至图11沿用原Windows实验'):
        set_text(p,'正文图1至图11沿用原Windows实验的真实截图；本附录图A1a、图A1b及图A2至图A9来自2026年10月21日Linux补充实验。终端图为实际窗口采集，浏览器图来自当次应用容器。图A1a在Word内仅裁去底部空白显示区，内嵌原PNG字节未改；所有关键命令与返回完整保留。')
    elif p.text.startswith('补充证据目录为docs/v3-evidence'):
        set_text(p,'本附录证据目录为docs/v3-evidence，结果文件为results.json。Windows原实验镜像为1.0.0，Linux补充镜像为1.0.1-v3；v5修订报告结构，应用JAR及这些历史实验结果不变。2026年10月22日本机复测另见附录B。各实验仓库均仅绑定本机。')
    elif p.text.startswith('本轮同时采集1440宽桌面和375宽手机页面'):
        set_text(p,'该轮同时采集1440宽桌面和375宽手机页面，手机页面未出现横向溢出；截图保留于证据包linux-v3目录。v5独立桌面及手机预览采用10月22日本机复测截图，两轮采集来源分别记录。')
    if re.match(r'^图(?:\d+|A\d+[ab]?)\s{2,}',p.text):
        set_text(p,re.sub(r'^(图(?:\d+|A\d+[ab]?))\s{2,}',r'\1 ',p.text))
    if p.text=='资源创建后的配置检查':
        p.style='Heading 3'
        for r in p.runs:fonts(r,'黑体','Times New Roman',12,True)

# Word crop hides only the blank lower canvas; the embedded evidence remains byte-identical.
shape=doc.inline_shapes[11]
fill=shape._inline.xpath('.//pic:blipFill')[0]
crop=OxmlElement('a:srcRect')
crop.set('b','41000')
fill.insert(1,crop)
shape.height=int(shape.height*.59)

task_map={
    'Spring Boot 项目打包发布':'Windows多阶段构建见第五章；Linux配套JAR及镜像见附录A2—A3；10月22日本机复测见附录B1',
    '发布后演示拉取':'Windows仓库localhost:5000见第六章；Linux仓库localhost:15085见附录A5—A8；本机15086见附录B1',
    '全过程截图汇总':'Windows图1—11；Linux图A1a、A1b及A2—A9；本机重建后图B1，均嵌入本Word',
    'Docker 环境部署与验证':'沿用已装引擎；Windows启动验证见2.2，Linux资源配置见附录A1；10月22日本机复测见B1',
}
for row in doc.tables[0].rows[1:]:
    if row.cells[0].text in task_map:set_text(row.cells[1].paragraphs[0],task_map[row.cells[0].text])

def add_p(text,style='Normal',page=False):
    p=doc.add_paragraph(style=style)
    p.paragraph_format.page_break_before=page
    p.paragraph_format.space_after=Pt(7)
    p.paragraph_format.line_spacing=Pt(12 if style=='Code' else (22 if style.startswith('Heading') else 16.5))
    p.paragraph_format.keep_with_next=style.startswith('Heading')
    p.paragraph_format.first_line_indent=Cm(.74 if style=='Normal' else 0)
    r=p.add_run(text)
    fonts(r,'黑体' if style.startswith('Heading') else '宋体','Consolas' if style=='Code' else 'Times New Roman',8.5 if style=='Code' else (18 if style=='Heading 1' else 14 if style=='Heading 2' else 11),style.startswith('Heading'))
    return p

add_p('附录B 本机复测与结果核对','Heading 1',True)
add_p('B1 2026年10月22日本机验证','Heading 2')
add_p('在原Windows机器恢复Docker引擎后，使用独立网络、容器和数据卷复测最新源码与交付包。Docker Client和Server均为29.8.2。本轮保留既有课堂项目及历史实验容器，未将测试成绩写入原作业数据库。')
table=doc.add_table(rows=1,cols=2)
table.style=doc.tables[0].style
table.autofit=False
table.columns[0].width=Cm(4)
table.columns[1].width=Cm(12.4)
for c,t in zip(table.rows[0].cells,['验证环节','本轮实际结果']):c.text=t
rows=[
('源码与容器构建','本机Maven测试3/3；Dockerfile多阶段构建成功，复用BuildKit缓存。容器内BOOT-INF的61项与交付JAR一致。'),
('浏览器与用户','实际容器Playwright 3/3；应用以garden普通用户运行，健康检查UP。'),
('独立资源','应用18087；Registry 15086；资源前缀memory-game6085-localtest-20261022，未复用原实验数据卷。'),
('镜像发布与拉取','标签localtest-20261022；push成功后移除本轮容器及两个本地标签；pull返回Downloaded newer image。'),
('新旧容器','91b23f11f6da → cdb872c45267；同一命名数据卷，完整排行榜JSON一致。'),
('真实测试成绩','测试花匠6085；easy；12步；9271 ms；记录1，仅为独立复测数据。'),
]
for row in rows:
    for c,t in zip(table.add_row().cells,row):c.text=t
for i,row in enumerate(table.rows):
    pr=row._tr.get_or_add_trPr()
    pr.append(OxmlElement('w:cantSplit'))
    if i==0:pr.append(OxmlElement('w:tblHeader'))
    for cell,width in zip(row.cells,[4,12.4]):
        cell.width=Cm(width)
        for p in cell.paragraphs:
            p.paragraph_format.first_line_indent=Cm(0)
            p.paragraph_format.line_spacing=Pt(14)
            p.paragraph_format.space_after=Pt(3)
            for r in p.runs:fonts(r,size=9,bold=i==0)

digest=next(x.split('@')[1] for x in RESULT['pullRepoDigests'] if x.startswith('localhost:15086/'))
add_p('本轮push/pull的同一仓库摘要','Normal')
add_p(digest,'Code')
add_p('仓库命令（项目目录内执行）','Normal')
add_p('docker push localhost:15086/memory-game6085:localtest-20261022\n'
      'docker pull localhost:15086/memory-game6085:localtest-20261022','Code')
add_p('本轮源码镜像JAR与交付JAR的整包摘要不同，但BOOT-INF内的应用类、前端资源和依赖逐项相同；未把整包摘要与应用内容校验混同。上述步骤的完整命令、删除检查及真实输出收录在证据包windows-localtest-20261022目录中。')

add_p('B2 拉取重建后的浏览器复核','Heading 2',True)
add_p('重新访问18087端口并切换到轻松排行榜，页面仍显示前一容器保存的“测试花匠6085”。浏览器无脚本错误，所见记录与重建前后相同的JSON响应一致。')
p=doc.add_paragraph()
p.paragraph_format.first_line_indent=Cm(0)
p.paragraph_format.line_spacing=1
p.paragraph_format.keep_with_next=True
p.alignment=WD_ALIGN_PARAGRAPH.CENTER
s=p.add_run().add_picture(str(LOCAL/'docker-roundtrip/after-recreate-ui.png'),width=Cm(16.4))
s._inline.docPr.set('descr','2026年10月22日从Registry重新拉取并重建后的真实页面及已保存排行榜')
p=add_p('图B1 本机拉取镜像重建后仍保留原测试成绩','Caption')
p.alignment=WD_ALIGN_PARAGRAPH.CENTER
add_p('证据来源：docs/local-test-20261022/docker-roundtrip/results.json及after-recreate-ui.png；构建与浏览器测试日志同收录证据包。应用和历史证据未随v5文档修订改写；本轮未重新安装Docker，也未补造原机器安装过程截图。')

# Classify intentional block formatting instead of changing source-code indentation.
for p in doc.paragraphs:
    t=p.text
    if not t:continue
    if p.style.name.startswith('Heading'):
        size=18 if p.style.name=='Heading 1' else 14 if p.style.name=='Heading 2' else 12
        for r in p.runs:fonts(r,'黑体','Times New Roman',size,True)
    elif re.match(r'^图(?:\d+|[AB]\d+[ab]?)\s',t):
        p.style='Caption'
        p.paragraph_format.first_line_indent=Cm(0)
        for r in p.runs:fonts(r,size=10)
    elif t.startswith('技术资料查阅日期：') or any(k in t for k in ['官方文档：https','官方仓库：https','官方部署文档：https']):
        p.style='Reference'
        p.paragraph_format.first_line_indent=Cm(0)
        ind=p._p.get_or_add_pPr().get_or_add_ind()
        for attr in ['firstLineChars','hangingChars','hanging']:
            ind.attrib.pop(qn('w:'+attr),None)
        for r in p.runs:fonts(r,size=10)
    elif (any('Mono' in (r.font.name or '') or (r.font.name or '')=='Consolas' for r in p.runs)
          or re.match(r'^(?:docker |curl\b|java |npm |\.\\scripts\\|Get-Item |FROM |RUN |COPY |WORKDIR |USER |ENV |EXPOSE |ENTRYPOINT |sha256:)',t)):
        p.style='Code'
        p.paragraph_format.first_line_indent=Cm(0)
        for r in p.runs:fonts(r,'宋体','Consolas',8.5)

# Reuse the established field-based TOC strategy with new unique bookmarks.
source=(ROOT/'Code/memory-game6085/scripts/build-v2-report.py').read_text(encoding='utf-8')
prefix=source.split('\ndoc = Document(SOURCE)',1)[0].replace('V2Heading','V5Heading').replace('enumerate(headings, 500)','enumerate(headings, 2000)')
namespace={'__file__':str(ROOT/'Code/memory-game6085/scripts/build-v2-report.py')}
exec(compile(prefix,'v5-toc-helper','exec'),namespace)
entries=namespace['build_toc'](doc)
for p in doc.paragraphs:
    if p._p.xpath('.//w:fldSimple[contains(@w:instr,"PAGEREF V5Heading")]'):
        p.style='toc 1' if p.paragraph_format.left_indent==Cm(0) else 'toc 2'
        p.paragraph_format.line_spacing=Pt(13)
        for r in p._p.xpath('.//w:r'):
            rp=r.find(qn('w:rPr'))
            if rp is not None:
                sz=rp.find(qn('w:sz'))
                if sz is not None:sz.set(qn('w:val'),'20')

doc.core_properties.revision=5
doc.core_properties.description='v5修复全部严格OOXML结构错误并补入本机Docker复测'
doc.save(DEST)

def tag_sequence(cls):
    for node in ast.parse(inspect.getsource(cls)).body[0].body:
        if isinstance(node,ast.Assign) and isinstance(node.targets[0],ast.Name) and node.targets[0].id=='_tag_seq':
            return ast.literal_eval(node.value)
    raise ValueError(cls)

orders={qn(tag):{qn(k):i for i,k in enumerate(tag_sequence(cls))} for tag,cls in [('w:pPr',CT_PPr),('w:rPr',CT_RPr),('w:tcPr',CT_TcPr),('w:settings',CT_Settings)]}
changes=Counter()
with ZipFile(DEST) as z:parts={n:z.read(n) for n in z.namelist()}
ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
for name,data in list(parts.items()):
    if not name.endswith('.xml'):continue
    root=E.fromstring(data)
    for shd in root.xpath('//w:shd[not(@w:val)]',namespaces=ns):
        shd.set(qn('w:val'),'clear')
        changes['missingShadingValue']+=1
    for node in root.iter():
        if node.tag not in orders:continue
        order=orders[node.tag]
        children=list(node)
        sorted_children=sorted(children,key=lambda c:order.get(c.tag,len(order)))
        if children!=sorted_children:
            node[:]=sorted_children
            changes[E.QName(node).localname+'Order']+=1
    parts[name]=E.tostring(root,encoding='UTF-8',xml_declaration=True,standalone=True)
with ZipFile(DEST,'w',ZIP_DEFLATED) as z:
    for name,data in parts.items():z.writestr(name,data)
(WORK/'toc-headings.json').write_text(json.dumps(entries,ensure_ascii=False,indent=2),encoding='utf-8')
(WORK/'normalization.json').write_text(json.dumps(dict(changes),ensure_ascii=False,indent=2),encoding='utf-8')
print(DEST)
print('Headings',len(entries),'Figures',len(doc.inline_shapes),'Tables',len(doc.tables),'Repairs',dict(changes))
