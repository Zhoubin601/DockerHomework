"""Build v4 from the verified report structure and unchanged experiment evidence."""
from pathlib import Path
from copy import deepcopy
import json
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[3]
WORK = ROOT / 'docs/v4-report-work'
EVIDENCE = ROOT / 'docs/v3-evidence'
WORK.mkdir(exist_ok=True)
RESULT = json.loads((EVIDENCE/'results.json').read_text())
doc = Document(ROOT/'output/v2/20261021-docker-lab-report-v2.docx')

def font(run, size=11, bold=False, mono=False):
    run.font.size, run.font.bold = Pt(size), bold
    run.font.color.rgb = RGBColor(0, 0, 0)
    rpr = run._r.get_or_add_rPr()
    fonts = rpr.find(qn('w:rFonts'))
    if fonts is None:
        fonts = OxmlElement('w:rFonts'); rpr.insert(0, fonts)
    fonts.attrib.clear()
    for attr in ['ascii', 'hAnsi', 'cs']:
        fonts.set(qn('w:'+attr), 'DejaVu Sans Mono' if mono else 'Liberation Serif')
    fonts.set(qn('w:eastAsia'), 'Noto Sans Mono CJK SC' if mono else 'Noto Serif CJK SC')

def paragraph(text, style='Normal', page=False, mono=False):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.page_break_before = page
    p.paragraph_format.first_line_indent = Cm(0 if mono or style.startswith('Heading') else .74)
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.line_spacing = Pt(12 if mono else (21 if style.startswith('Heading') else 16.5))
    p.paragraph_format.keep_with_next = style.startswith('Heading')
    font(p.add_run(text), 8.5 if mono else (14 if style.startswith('Heading') else 11), style.startswith('Heading'), mono)
    return p

def replace(p, text):
    keep = deepcopy(p.runs[0]._r.rPr) if p.runs and p.runs[0]._r.rPr is not None else None
    p.clear(); run = p.add_run(text)
    if keep is not None: run._r.insert(0, keep)

def insert_after(anchor, text):
    p = paragraph(text)
    anchor._p.addnext(p._p)
    return p

def picture(name, caption):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.line_spacing = 1
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_after = Pt(2)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    shape = p.add_run().add_picture(str(EVIDENCE/name), width=Cm(16.4))
    shape._inline.docPr.set('descr', caption + ' 实际终端或浏览器截图')
    p = paragraph(caption)
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.line_spacing = Pt(14)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in p.runs: font(run, 10)

for node in doc.element.xpath('.//w:lastRenderedPageBreak'):
    node.getparent().remove(node)
for p in list(doc.paragraphs):
    if p.text.startswith('本实验将 Vue3'):
        replace(p, '本实验将 Vue3、Spring Boot 与 SQLite 记忆翻牌小游戏打包为 Docker 镜像，完成本机仓库发布、删除本地镜像后拉取和数据卷保留验证。正文记录原 Windows 实验，附录记录独立 Linux 云工作区的补充实验。v4 补齐已有资源创建截图和关键命令；两次实验的镜像版本、端口与结果分别保留。')
    elif p.text == '二 实验环境与项目结构':
        insert_after(p, '第二至九章记录原 Windows 实验。v3 的 Linux 补充部署使用独立资源，环境配置和实际操作截图见附录A。')
    elif p.text.startswith('引擎启动后，在项目根目录执行版本与命名检查'):
        insert_after(p, '原 Windows 环境沿用既有安装。本报告附录A补充本轮 Linux 环境的网络、数据卷、镜像及容器配置操作截图，两次环境的记录分别保留。')
    elif p.text == '9.1 截图来源索引':
        replace(p, '9.1 原Windows实验截图来源索引')

paragraph('附录A Linux补充部署与复现', 'Heading 1', page=True)
paragraph('A1 环境来源与独立资源配置', 'Heading 2')
paragraph('2026年10月21日在 Debian GNU/Linux 13 云工作区补做部署。平台已提供 Docker Engine，本轮沿用其本地守护进程；Client 与 Server 均为28.4.0，Compose为2.40.3，存储驱动为vfs。')
paragraph('实际创建独立网络 memory-game6085-v3-net，以及应用和仓库数据卷。图A1a记录首次创建命令及返回结果，图A1b记录随后配置检查；完整日志保存在证据包。两次实验分别使用独立资源。')
picture('01-resource-creation.png', '图A1a 首次创建独立网络及数据卷')
continuation = paragraph('资源创建后的配置检查', page=True)
continuation.paragraph_format.first_line_indent = Cm(0)
continuation.paragraph_format.keep_with_next = True
for run in continuation.runs:
    font(run, 14, True)
paragraph('创建资源后，分别检查引擎与Compose版本、网络驱动和应用及仓库数据卷。检查返回network为bridge、volume为local，与首次创建结果对应。')
picture('01-config-review.png', '图A1b Linux引擎及网络和数据卷配置验证')
paragraph('两轮实验均沿用已安装的Docker引擎。本附录记录当前环境的资源配置与应用部署；本报告未附原Windows安装、首次配置和启动恢复的历史操作截图，现有记录不包含引擎安装全过程。')

paragraph('A2 配套JAR打包为新镜像', 'Heading 2', page=True)
paragraph('使用 Dockerfile.prebuilt 将配套可执行 JAR 复制到 Temurin 21 JRE 镜像，配置普通用户garden和数据目录/app/data。此步骤使用已完成 Maven 打包的交付 JAR；源码多阶段构建流程仍见第五章。')
paragraph('交付JAR的SHA256\n' + RESULT['jarSha256'], mono=True)
paragraph('docker build --progress=plain --pull=false -f Dockerfile.prebuilt -t memory-game6085:1.0.1-v3 .', mono=True)
picture('02-build-review.png', '图A2 校验交付JAR并成功构建1.0.1-v3镜像')

paragraph('A3 容器运行与交付文件校验', 'Heading 2', page=True)
paragraph('将镜像运行在独立网络中，宿主机127.0.0.1:18086映射容器8080，并挂载memory-game6085-v3-data。健康接口返回UP，进程身份为garden。')
paragraph('分别计算交付JAR与容器/app/app.jar的SHA256，两个结果完全一致。这将本轮交付文件与实际部署镜像直接对应。')
picture('03-container-run.png', '图A3 新容器运行成功且容器内外JAR摘要相同')

score = RESULT['savedScore']
paragraph('A4 实际游戏与成绩保存', 'Heading 2', page=True)
paragraph(f'通过浏览器在本轮Linux容器完成轻松难度牌局，并提交昵称“{score["nickname"]}”。实际记录编号为{score["id"]}，{score["moves"]}步，用时{score["durationMs"]} ms；数据由服务端生成并保存到SQLite。')
paragraph('本轮同时采集1440宽桌面和375宽手机页面，手机页面未出现横向溢出。配套预览图片来自本次容器访问。')
picture('05-game-saved.png', '图A4 Linux容器内完成牌局并实际保存成绩')

paragraph('A5 本机仓库发布镜像', 'Heading 2', page=True)
paragraph('启动独立Registry容器，绑定127.0.0.1:15085，并使用专用仓库数据卷。将新镜像标记到localhost:15085后执行push，仓库标签查询返回1.0.1-v3。此地址仅用于本机实验。')
picture('04-registry-push.png', '图A5 启动本机Registry并实际推送新镜像')

paragraph('A6 移除本地镜像并保留数据卷', 'Heading 2', page=True)
paragraph('在发布成功并记录完整排行榜后，移除本轮应用容器和两个1.0.1-v3本地标签。按版本过滤的镜像列表为空；命名数据卷仍存在。移除操作仅针对本轮独立资源。')
picture('07-remove-local.png', '图A6 应用容器及本地镜像已移除而数据卷保留')

paragraph('A7 从仓库重新拉取', 'Heading 2', page=True)
paragraph('本地版本移除后执行docker pull，结果为Downloaded newer image。推送与拉取的仓库摘要一致，与原Windows实验的摘要分别记录。')
paragraph('Linux补充实验仓库摘要\n' + RESULT['repoDigests'][0].split('@')[1], mono=True)
picture('08-registry-pull.png', '图A7 本地镜像移除后从Registry实际拉取成功')

paragraph('A8 重建容器并核对成绩保留', 'Heading 2', page=True)
paragraph('使用刚拉取的仓库镜像创建新容器，并挂载相同命名数据卷。原容器前12位为' + RESULT['oldContainer'][:12] + '，新容器前12位为' + RESULT['newContainer'][:12] + '；新旧标识不同。')
paragraph('健康检查再次通过，容器内JAR仍与交付文件一致，重建前后的完整排行榜JSON相同。比较结果为LeaderboardPreserved=True、ContainerChanged=True。')
picture('09-recreate-verify.png', '图A8 拉取镜像重建后JAR及完整排行榜校验通过')

paragraph('A9 浏览器复核重建后的榜单', 'Heading 2', page=True)
paragraph('重新访问本轮应用端口并切换到轻松难度，页面仍显示先前保存的真实成绩。该页面由新容器提供，与终端JSON比较结果一致。')
picture('10-after-recreate-ui.png', '图A9 新容器页面仍显示原已保存成绩')

paragraph('A10 补充证据来源与版本对应', 'Heading 2', page=True)
paragraph('正文图1至图11沿用原Windows实验的真实截图；本附录图A1a、图A1b及图A2至图A9来自2026年10月21日Linux补充实验。终端图是实际终端窗口的像素采集，浏览器图来自当次应用容器。v4新增图A1a并将关键命令整理为可复制文字，原始截图与日志保持原字节。')
table = doc.add_table(rows=1, cols=3)
table.style = doc.tables[0].style
table.autofit = False
for col, width in zip(table.columns, [1.3, 4.8, 10.3]): col.width = Cm(width)
for c, t in zip(table.rows[0].cells, ['图号', '内容', '来源文件']): c.text = t
sources = [
('A1a','首次资源创建','01-resource-creation.png'),('A1b','配置检查','01-config-review.png'),('A2','JAR与镜像构建','02-build-review.png'),
('A3','运行与JAR一致','03-container-run.png'),('A4','游戏成绩保存','05-game-saved.png'),
('A5','本机仓库发布','04-registry-push.png'),('A6','本地镜像移除','07-remove-local.png'),
('A7','仓库拉取','08-registry-pull.png'),('A8','重建与数据比较','09-recreate-verify.png'),
('A9','重建后页面','10-after-recreate-ui.png')]
for row in sources:
    for c,t in zip(table.add_row().cells,row): c.text=t
for index,row in enumerate(table.rows):
    trpr=row._tr.get_or_add_trPr()
    trpr.append(OxmlElement('w:cantSplit'))
    if index == 0: trpr.append(OxmlElement('w:tblHeader'))
    for c,width in zip(row.cells,[1.3,4.8,10.3]):
        c.width=Cm(width)
        tcpr=c._tc.get_or_add_tcPr()
        if index == 0:
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E8EEEB');tcpr.append(shade)
        borders=OxmlElement('w:tcBorders')
        for edge in ['top','left','bottom','right']:
            border=OxmlElement('w:'+edge)
            for attr,value in [('val','single'),('sz','4'),('color','D9D9D9')]: border.set(qn('w:'+attr),value)
            borders.append(border)
        tcpr.append(borders)
        for p in c.paragraphs:
            p.paragraph_format.first_line_indent=Cm(0)
            p.paragraph_format.line_spacing=Pt(14)
            for run in p.runs: font(run,9,index==0)
paragraph('补充证据目录为docs/v3-evidence，结果文件为results.json。Windows原实验镜像为1.0.0，Linux补充镜像为1.0.1-v3；v4为报告与交付资料修订，未更改应用JAR、镜像标签或实验结果。两套仓库均仅绑定本机，复现实验时应在实际执行机器上使用对应地址和数据卷。')

# Update the acceptance map without altering historical measurements.
task_map = {
    'Spring Boot 项目打包发布': 'Windows多阶段构建见第五章；Linux配套JAR及镜像见附录A2—A3',
    '发布后演示拉取': 'Windows仓库localhost:5000见第六章；Linux仓库localhost:15085见附录A5—A8',
    '全过程截图汇总': 'Windows图1—11；Linux图A1a、A1b及A2—A9，均嵌入本Word',
    'Docker 环境部署与验证': '沿用已装引擎；Windows启动验证见2.2，Linux资源创建及检查见附录A1',
}
for row in doc.tables[0].rows[1:]:
    if row.cells[0].text in task_map:
        replace(row.cells[1].paragraphs[0], task_map[row.cells[0].text])
        for p in row.cells[1].paragraphs:
            p.paragraph_format.line_spacing = Pt(14)

# Commands below are transcribed from unchanged original logs, not new runs.
def before_figure(caption, text, mono=False):
    target = next(p for p in doc.paragraphs if p.text == caption)
    image = target._p.getprevious()
    while image is not None and not image.xpath('.//w:drawing'):
        image = image.getprevious()
    assert image is not None, caption
    p = paragraph(text, mono=mono)
    image.addprevious(p._p)
    return p

jar_hash = RESULT['jarSha256']
digest = RESULT['repoDigests'][0].split('@')[1]
before_figure('图A1a 首次创建独立网络及数据卷',
    'docker network create --label course=memory-game6085-v3 \\\n'
    '  memory-game6085-v3-net\n'
    'docker volume create --label course=memory-game6085-v3 \\\n'
    '  memory-game6085-v3-data\n'
    'docker volume create --label course=memory-game6085-v3 \\\n'
    '  memory-game6085-v3-registry-data', mono=True)
before_figure('图A3 新容器运行成功且容器内外JAR摘要相同',
    '关键命令摘自03-container-run.txt。以下文本便于复制，实际返回结果见原图。')
before_figure('图A3 新容器运行成功且容器内外JAR摘要相同',
    'docker run -d --name memory-game6085-v3-app \\\n'
    '  --network memory-game6085-v3-net \\\n'
    '  -p 127.0.0.1:18086:8080 -e DATA_DIR=/app/data \\\n'
    '  --mount type=volume,source=memory-game6085-v3-data,target=/app/data \\\n'
    '  memory-game6085:1.0.1-v3\n'
    'curl --fail --silent http://127.0.0.1:18086/api/health\n'
    'docker exec memory-game6085-v3-app sha256sum /app/app.jar', mono=True)
before_figure('图A3 新容器运行成功且容器内外JAR摘要相同',
    '交付JAR与容器内文件SHA256相同\n' + jar_hash, mono=True)
before_figure('图A5 启动本机Registry并实际推送新镜像',
    '关键命令及结果摘自04-registry-push.txt。')
before_figure('图A5 启动本机Registry并实际推送新镜像',
    'docker tag memory-game6085:1.0.1-v3 \\\n'
    '  localhost:15085/memory-game6085:1.0.1-v3\n'
    'docker push localhost:15085/memory-game6085:1.0.1-v3\n'
    'curl --fail --silent \\\n'
    '  http://127.0.0.1:15085/v2/memory-game6085/tags/list', mono=True)
before_figure('图A5 启动本机Registry并实际推送新镜像',
    '推送与后续拉取的仓库摘要一致\n' + digest, mono=True)
before_figure('图A8 拉取镜像重建后JAR及完整排行榜校验通过',
    '关键命令及结果摘自09-recreate-verify.txt；重建继续使用相同数据卷。')
before_figure('图A8 拉取镜像重建后JAR及完整排行榜校验通过',
    'docker run -d --name memory-game6085-v3-pulled \\\n'
    '  --network memory-game6085-v3-net \\\n'
    '  -p 127.0.0.1:18086:8080 -e DATA_DIR=/app/data \\\n'
    '  --mount type=volume,source=memory-game6085-v3-data,target=/app/data \\\n'
    '  localhost:15085/memory-game6085:1.0.1-v3\n'
    'docker exec memory-game6085-v3-pulled sha256sum /app/app.jar', mono=True)
before_figure('图A8 拉取镜像重建后JAR及完整排行榜校验通过',
    '容器内JAR SHA256\n' + jar_hash + '\n'
    'LeaderboardPreserved=True ContainerChanged=True', mono=True)

# Rebuild the original TOC using its existing formatting strategy and new bookmarks.
source=(ROOT/'Code/memory-game6085/scripts/build-v2-report.py').read_text()
prefix=source.split('\ndoc = Document(SOURCE)',1)[0].replace('V2Heading','V4Heading').replace('enumerate(headings, 500)','enumerate(headings, 1000)')
namespace={'__file__':str(ROOT/'Code/memory-game6085/scripts/build-v2-report.py')}
exec(compile(prefix,'v4-toc-helper','exec'),namespace)
entries=namespace['build_toc'](doc)
for p in doc.paragraphs:
    if p._p.xpath('.//w:fldSimple[contains(@w:instr,"PAGEREF V4Heading")]'):
        p.paragraph_format.line_spacing=Pt(14)
        for run in p._p.xpath('.//w:r'):
            rpr=run.find(qn('w:rPr'))
            if rpr is None: rpr=OxmlElement('w:rPr');run.insert(0,rpr)
            fonts=rpr.find(qn('w:rFonts'))
            if fonts is None: fonts=OxmlElement('w:rFonts');rpr.insert(0,fonts)
            fonts.attrib.clear()
            for a in ['ascii','hAnsi','cs']:fonts.set(qn('w:'+a),'Liberation Serif')
            fonts.set(qn('w:eastAsia'),'Noto Serif CJK SC')
doc.core_properties.revision=4
doc.core_properties.description='v4补齐资源创建图及可复制命令，原Windows与Linux实验分别记录'
dest=WORK/'20261021-docker-lab-report-v4.docx'
doc.save(dest)
(WORK/'toc-headings.json').write_text(json.dumps(entries,ensure_ascii=False,indent=2))
(WORK/'supplement-figures.json').write_text(json.dumps(sources,ensure_ascii=False,indent=2))
print(dest)
print('TOC',len(entries),'Tables',len(doc.tables),'Figures',len(doc.inline_shapes))
